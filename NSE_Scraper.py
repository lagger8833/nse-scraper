import time
import os
from typing import List, Tuple
import yfinance as yf
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference

# NIFTY 50 tickers list
NIFTY_50_TICKERS = [
    "RELIANCE.NS", "TCS.NS", "HDFCBANK.NS", "BHARTIARTL.NS", "ICICIBANK.NS", 
    "HINDUNILVR.NS", "INFY.NS", "SBIN.NS", "KOTAKBANK.NS", "ITC.NS", 
    "SUNPHARMA.NS", "LT.NS", "BAJFINANCE.NS", "HCLTECH.NS", "MARUTI.NS", 
    "NTPC.NS", "ULTRACEMCO.NS", "AXISBANK.NS", "M&M.NS", "BAJAJFINSV.NS",
    "ONGC.NS",  "TITAN.NS", "POWERGRID.NS", "ADANIPORTS.NS", "ADANIENT.NS",
    "WIPRO.NS", "JSWSTEEL.NS", "ASIANPAINT.NS", "COALINDIA.NS", "NESTLEIND.NS", 
    "TATAMOTORS.NS", "BAJAJ-AUTO.NS", "GRASIM.NS", "TRENT.NS", "SBILIFE.NS", 
    "TATASTEEL.NS", "EICHERMOT.NS", "HDFCLIFE.NS", "ZOMATO.NS", "BEL.NS", 
    "HEROMOTOCO.NS","TECHM.NS", "HINDALCO.NS", "SHRIRAMFIN.NS", "TATACONSUM.NS",
    "APOLLOHOSP.NS", "DRREDDY.NS", "CIPLA.NS", "JIOFIN.NS", "INDUSINDBK.NS"
]

EXCEL_FILE = "nifty50_latest_snapshot.xlsx"


def fetch_stock_data(ticker: str) -> Tuple[float, float]: 
    try:
        stock = yf.Ticker(ticker)
        data = stock.history(period="1d")

        if data.empty or 'Open' not in data.columns or 'Close' not in data.columns:
            print(f"No valid data found for {ticker}. It might be delisted or market is closed.")
            return None, None

        open_price = data['Open'].iloc[0]
        close_price = data['Close'].iloc[0]
        return round(open_price, 2), round(close_price, 2)

    except Exception as e:
        print(f"Error fetching data for {ticker}: {e}")
        return None, None


def write_current_snapshot_with_chart(stock_data: List[Tuple[str, float, float, float]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock Data"

    # Write header
    ws.append(["Stock", "Open", "Current", "Change (%)"])

    # Write stock data
    for row in stock_data:
        ws.append(row)

    # Calculate average percentage change
    avg_change = round(sum(row[3] for row in stock_data) / len(stock_data), 2)

    # Append a row with the average
    ws.append(["Average", None, None, avg_change])

    # Add chart (excluding the Average row)
    chart = BarChart()
    chart.title = "NIFTY 50 - % Change"
    chart.y_axis.title = "Change (%)"
    chart.x_axis.title = "Stock"

    chart.width = 20
    chart.height = 12

    num_rows = len(stock_data)
    data = Reference(ws, min_col=4, min_row=2, max_row=num_rows+1)  # exclude average
    cats = Reference(ws, min_col=1, min_row=2, max_row=num_rows+1)

    chart.add_data(data, titles_from_data=False)
    chart.set_categories(cats)

    ws.add_chart(chart, f"F2")

    # Save Excel file
    wb.save(EXCEL_FILE)


if __name__ == "__main__":
    print("Tracking NIFTY 50 stock data every 1 minute. Press Ctrl+C to stop.")
    
    try:
        while True:
            all_stock_data : List[Tuple[str, float, float, float]] = []

            for ticker in NIFTY_50_TICKERS:
                open_price, current_price = fetch_stock_data(ticker)
                if open_price is None or current_price is None:
                    continue
                change_pct = ((current_price - open_price) / open_price) * 100
                all_stock_data.append([
                    ticker.replace('.NS', ''),
                    open_price,
                    current_price,
                    round(change_pct, 2)
                ])

            write_current_snapshot_with_chart(all_stock_data)

            # Show in terminal
            df = pd.DataFrame(all_stock_data, columns=["Stock", "Open", "Current", "Change (%)"])
            print(df)

            time.sleep(60)
    except KeyboardInterrupt:
        print("Stopped by user.")
